"""Generate DOCX and PDF from SaaS_Salon_System_Documentation.md"""
import re
import sys
from pathlib import Path

DOCS_DIR = Path(__file__).parent
MD_PATH = DOCS_DIR / "SaaS_Salon_System_Documentation.md"
DOCX_PATH = DOCS_DIR / "SaaS_Salon_System_Documentation.docx"
PDF_PATH = DOCS_DIR / "SaaS_Salon_System_Documentation.pdf"
ROOT = DOCS_DIR.parent


def generate_docx():
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_buffer: list[str] = []
    in_table = False
    table_rows: list[list[str]] = []

    def flush_code():
        nonlocal code_buffer, in_code
        if not code_buffer:
            return
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_buffer))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        code_buffer = []
        in_code = False

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(table_rows):
            for ci, cell in enumerate(row):
                t.rows[ri].cells[ci].text = cell.strip()
        doc.add_paragraph()
        table_rows = []
        in_table = False

    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                flush_code()
            else:
                flush_table()
                in_code = True
            continue

        if in_code:
            code_buffer.append(line)
            continue

        if line.strip().startswith("|") and "|" in line[1:]:
            if re.match(r"^\|[-:\s|]+\|$", line.strip()):
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_rows.append(cells)
            in_table = True
            continue
        elif in_table:
            flush_table()

        if line.startswith("# ") and not line.startswith("##"):
            flush_table()
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            flush_table()
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            flush_table()
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            flush_table()
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == "---":
            doc.add_paragraph("─" * 40)
        elif line.strip().startswith(">"):
            p = doc.add_paragraph(line.strip().lstrip("> ").strip())
            p.paragraph_format.left_indent = Inches(0.25)
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif line.strip() == "":
            pass
        else:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            text = re.sub(r"`(.+?)`", r"\1", text)
            text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
            if text.strip():
                doc.add_paragraph(text.strip())

    flush_code()
    flush_table()
    doc.save(DOCX_PATH)
    print(f"Created: {DOCX_PATH}")


def generate_pdf():
    try:
        from fpdf import FPDF
    except ImportError:
        print("fpdf2 not installed; skipping PDF", file=sys.stderr)
        return False

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 8, "SaaS Salon - System Documentation", align="C", new_x="LMARGIN", new_y="NEXT")

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 10, f"Page {self.page_no()}", align="C")

        def write_wrapped(self, text: str, h: float = 5):
            text = safe_text(text)
            if not text.strip():
                return
            w = self.epw
            # Break unbreakable long tokens (URLs, paths)
            parts = re.split(r"(\s+)", text)
            for part in parts:
                if not part:
                    continue
                if len(part) <= 100:
                    self.multi_cell(w, h, part)
                    continue
                for i in range(0, len(part), 80):
                    self.multi_cell(w, h, part[i : i + 80])

    def safe_text(s: str) -> str:
        s = s.replace("\t", "    ")
        s = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", "?", s)
        return s

    pdf = PDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)

    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    in_code = False
    skip_mermaid = False

    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                in_code = False
                skip_mermaid = False
            else:
                in_code = True
                skip_mermaid = line.strip().startswith("```mermaid")
            continue
        if in_code:
            if skip_mermaid:
                continue
            pdf.set_font("Courier", size=8)
            pdf.write_wrapped(line, h=4)
            pdf.set_font("Helvetica", size=10)
            continue

        if line.startswith("# ") and not line.startswith("##"):
            pdf.ln(4)
            pdf.set_font("Helvetica", "B", 16)
            pdf.write_wrapped(line[2:].strip(), h=8)
            pdf.set_font("Helvetica", size=10)
        elif line.startswith("## "):
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 13)
            pdf.write_wrapped(line[3:].strip(), h=7)
            pdf.set_font("Helvetica", size=10)
        elif line.startswith("### "):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.write_wrapped(line[4:].strip(), h=6)
            pdf.set_font("Helvetica", size=10)
        elif line.strip().startswith("|"):
            if re.match(r"^\|[-:\s|]+\|$", line.strip()):
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            pdf.write_wrapped(" | ".join(cells))
        elif line.strip().startswith("- "):
            pdf.write_wrapped("  - " + line.strip()[2:])
        elif line.strip() == "---":
            pdf.ln(2)
        elif line.strip().startswith(">"):
            pdf.set_font("Helvetica", "I", 10)
            pdf.write_wrapped(line.strip().lstrip("> "))
            pdf.set_font("Helvetica", size=10)
        elif line.strip():
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            text = re.sub(r"`(.+?)`", r"\1", text)
            text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
            pdf.write_wrapped(text.strip())

    pdf.output(PDF_PATH)
    print(f"Created: {PDF_PATH}")
    return True


def generate_pdf_from_docx():
    try:
        from docx2pdf import convert

        convert(str(DOCX_PATH), str(PDF_PATH))
        print(f"Created: {PDF_PATH} (via Word)")
        return True
    except Exception as exc:
        print(f"docx2pdf failed ({exc}); falling back to fpdf", file=sys.stderr)
        return generate_pdf()


if __name__ == "__main__":
    generate_docx()
    generate_pdf_from_docx()
    for path in (DOCX_PATH, PDF_PATH, MD_PATH):
        if path.exists():
            import shutil
            shutil.copy2(path, ROOT / path.name)
            print(f"Copied: {ROOT / path.name}")
